#!/usr/bin/env python3
"""presentation-studio deck kit (Python / python-pptx) — same deck.json schema as the JS kit.

Usage:  python deck_kit.py deck.json [out.pptx]
Needs:  python -m pip install python-pptx   (Pillow optional, for image ratios)
Exit code 2 when there are warnings (fix them, rebuild). Speaker notes use the agreed template:
[i/N] Hedef/Bitişte header · say · ÖZETLE · ARAŞTIRMA (neden/ne bulundu/nasıl yorumlanır) · TERİMLER · → Geçiş
"""
from __future__ import annotations

import json
import math
import os
import re
import sys

try:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt
except ImportError:
    sys.exit("python-pptx is missing: python -m pip install python-pptx")

SAFE_FONTS = ["Calibri", "Arial", "Cambria", "Times New Roman", "Georgia", "Verdana", "Tahoma"]
FONT_EM = {"Calibri": 0.46, "Arial": 0.5, "Verdana": 0.56, "Tahoma": 0.5, "Cambria": 0.5, "Georgia": 0.52, "Times New Roman": 0.45}
THEMES = {
    "klinik": dict(dark="15303D", ink="1E2A33", muted="5A6B76", light="F3F6F8", card="E8EEF2", accent="0E7C86", accent2="D2553F", accent3="E9A23B", white="FFFFFF"),
    "akademik": dict(dark="1B2430", ink="1F2733", muted="5B6676", light="F4F5F8", card="E9ECF2", accent="B0303F", accent2="0F6E8C", accent3="E0A43A", white="FFFFFF"),
    "sade": dict(dark="22262E", ink="23272F", muted="626A75", light="F5F6F7", card="ECEEF1", accent="2F5DA8", accent2="C2453A", accent3="D99A2B", white="FFFFFF"),
    "doga": dict(dark="1F3326", ink="1F2A22", muted="5C6B60", light="F4F6F2", card="E7ECE4", accent="3C7A4A", accent2="B8573A", accent3="D6A33A", white="FFFFFF"),
}
DEFAULT_TYPE = dict(minBody=36, title=40, body=36, stat=80, statLabel=36, section=54, sectionNum=110, coverTitle=54, coverSub=36,
                    exceptions=dict(source=16, references=16, slideNumber=14, credit=14))
W, H, MX = 13.333, 7.5, 0.6
G = dict(titleY=0.42, titleH=1.25, bodyY=1.85, bodyB=6.75, sourceY=6.95, sourceH=0.38)
NBSP = " "


def rgb(h):
    return RGBColor.from_string(h)


def no_break(t: str) -> str:
    t = re.sub(r"(\d) (?=(%|N·m/kg|N·m|Nm|N|kg|kgf|g|ms|sn|s|dk|min|m|cm|mm|°|yaş|kişi|hasta|hafta|gün|Hz|dB|mmHg|bpm)(?!\w))", r"\1" + NBSP, str(t))
    t = re.sub(r"([≥≤≈<>=±→×÷]) (?=[\d%−\-~])", r"\1" + NBSP, t)
    return t


def plain(s) -> str:
    return re.sub(r"\*\*|\^\{|_\{|\}", "", str(s or ""))


class Deck:
    def __init__(self, meta: dict, base_dir: str):
        self.meta = meta
        self.base = base_dir
        self.t = {**THEMES.get(meta.get("theme", "klinik"), THEMES["klinik"]), **meta.get("colors", {})}
        ty = meta.get("type", {})
        self.type = {**DEFAULT_TYPE, **ty, "exceptions": {**DEFAULT_TYPE["exceptions"], **ty.get("exceptions", {})}}
        self.font = meta.get("font", "Calibri")
        self.em = FONT_EM.get(self.font, 0.5)
        if self.font not in SAFE_FONTS:
            print(f"[kit] font {self.font} not in safe list {SAFE_FONTS} — ask the user before shipping")
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.blank = self.prs.slide_layouts[6]
        self.warnings: list[str] = []
        self.n = 0
        self.log: list[str] = []

    # ------------------------------------------------------------ measuring
    def lines(self, text, fs, w):
        per = max(1, int(w / (fs * self.em / 72)))
        return sum(max(1, math.ceil(len(ln) / per)) for ln in plain(text).split("\n"))

    def block_h(self, text, fs, w, sp=1.2):
        return self.lines(text, fs, w) * fs * sp / 72

    def floor(self, role, fs=None):
        ex = self.type["exceptions"]
        if role in ex and ex[role] is not None:
            return fs if fs is not None else ex[role]
        return max(fs if fs is not None else self.type["body"], self.type["minBody"])

    # ------------------------------------------------------------ primitives
    def _runs(self, para, text, color, bold=False, size=None):
        for part in re.split(r"(\*\*[^*]+\*\*|\^\{[^}]+\}|_\{[^}]+\})", no_break(text)):
            if not part:
                continue
            r = para.add_run()
            if part.startswith("**"):
                r.text, b = part[2:-2], True
            elif part.startswith("^{"):
                r.text, b = part[2:-1], bold
                r.font._rPr.set("baseline", "30000")
            elif part.startswith("_{"):
                r.text, b = part[2:-1], bold
                r.font._rPr.set("baseline", "-25000")
            else:
                r.text, b = part, bold
            r.font.bold = b
            r.font.size = Pt(size)
            r.font.name = self.font
            r.font.color.rgb = rgb(color)

    def text(self, slide, s, x, y, w, h, role="body", fs=None, bold=False, color=None, align="left", valign="top", name=None, spacing=1.2):
        fs = self.floor(role, fs)
        if self.block_h(s, fs, w, spacing) > h + 0.02:
            self.warnings.append(f'slide {self.n}: "{plain(s)[:40]}…" may overflow ({fs} pt in {w:.2f}×{h:.2f} in)')
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tb.name = name or f"Metin:{role}"
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[valign]
        for i, line in enumerate(str(s).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
            self._runs(p, line, color or self.t["ink"], bold, fs)
        return tb

    def bullets(self, slide, items, x, y, w, h, color=None):
        fs = self.floor("body")
        need = sum(self.block_h(it, fs, w - 0.45, 1.2) for it in items) + (len(items) - 1) * 12 / 72
        if need > h + 0.05:
            self.warnings.append(f"slide {self.n}: bullet list may overflow ({need:.2f} of {h:.2f} in) — shorten or split")
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tb.name = "Metin:maddeler"
        tf = tb.text_frame
        tf.word_wrap = True
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Emu(Inches(0.42))))
            pPr.set("indent", str(-Emu(Inches(0.42))))
            bu = pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar", {"char": "•"})
            pPr.append(bu)
            if i:
                p.space_before = Pt(12)
            self._runs(p, it, color or self.t["ink"], False, fs)

    def rect(self, slide, x, y, w, h, fill, name="Kutu", rounded=True):
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.name = name
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
        sh.line.fill.background()
        sh.shadow.inherit = False
        if rounded:
            sh.adjustments[0] = 0.06
        return sh

    def chip(self, slide, label, x, y, d, fill=None):
        sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
        sh.name = "Numara"
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill or self.t["accent"])
        sh.line.fill.background()
        sh.shadow.inherit = False
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        self._runs(p, str(label), self.t["white"], True, max(self.type["minBody"], round(d * 38)))

    def image(self, slide, img, x, y, w, h, mode="contain"):
        f = img["file"] if os.path.isabs(img["file"]) else os.path.join(self.base, img["file"])
        if not os.path.exists(f):
            self.warnings.append(f"slide {self.n}: image missing {img['file']} — placeholder drawn")
            self.rect(slide, x, y, w, h, self.t["card"], "Gorsel yok")
            self.text(slide, "Görsel: " + (img.get("alt") or img["file"]), x + 0.2, y + h / 2 - 0.3, w - 0.4, 0.6, role="credit", fs=18, color=self.t["muted"], align="center")
            return
        r = img.get("ratio")
        if not r:
            try:
                from PIL import Image
                with Image.open(f) as im:
                    r = im.width / im.height
            except Exception:  # noqa: BLE001
                r = 1.5
        if mode == "cover":
            pic = slide.shapes.add_picture(f, Inches(x), Inches(y), Inches(w), Inches(h))
            box = w / h
            if r > box:  # wider: crop left/right
                c = (1 - box / r) / 2
                pic.crop_left = pic.crop_right = c
            else:
                c = (1 - r / box) / 2
                pic.crop_top = pic.crop_bottom = c
        else:
            iw, ih = w, w / r
            if ih > h:
                ih, iw = h, h * r
            pic = slide.shapes.add_picture(f, Inches(x + (w - iw) / 2), Inches(y + (h - ih) / 2), Inches(iw), Inches(ih))
        pic.name = "Gorsel"
        if img.get("alt"):
            pic._element.nvPicPr.cNvPr.set("descr", img["alt"])

    def bg(self, slide, color):
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(color)

    def frame(self, slide, s, dark=False, number=True, title=True):
        t = self.t
        if s.get("kicker") and not self.type["exceptions"].get("kicker"):
            self.warnings.append(f"slide {self.n}: kicker ignored — small labels are off unless exceptions.kicker is set")
        if s.get("kicker") and self.type["exceptions"].get("kicker"):
            self.text(slide, s["kicker"].upper(), MX, 0.22, W - 2 * MX, 0.36, role="kicker", bold=True, color=t["accent3"] if dark else t["accent"], name="Ust etiket")
        if title and s.get("title"):
            self.text(slide, s["title"], MX, G["titleY"] + (0.1 if s.get("kicker") else 0), W - 2 * MX, G["titleH"], role="title", fs=self.type["title"], bold=True, color=t["white"] if dark else t["dark"], name="Baslik")
        mut = "C9D2DA" if dark else t["muted"]
        if s.get("source"):
            self.text(slide, s.get("sourcePrefix", "Kaynak: ") + s["source"], MX, G["sourceY"], W - 2 * MX - 1.0, G["sourceH"], role="source", color=mut, name="Kaynak satiri", valign="middle")
        if s.get("credit"):
            cy = G["sourceY"] + 0.05 if (not number and not s.get("source")) else G["sourceY"] - 0.32
            self.text(slide, s["credit"], MX, cy, W - 2 * MX - 1.0, 0.3, role="credit", color=mut, name="Gorsel kredisi")
        if number:
            self.text(slide, str(self.n), W - MX - 0.8, G["sourceY"], 0.8, G["sourceH"], role="slideNumber", color=mut, align="right", name="Slayt no", valign="middle")
        m = slide.shapes.add_textbox(0, 0, Inches(0.01), Inches(0.01))
        m.name = "layout:" + s["layout"]
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]

    # ------------------------------------------------------------ layouts
    def new(self, color):
        sl = self.prs.slides.add_slide(self.blank)
        self.bg(sl, color)
        return sl

    def cover(self, s):
        t, sl = self.t, self.new(self.t["dark"])
        tw = 7.0 if s.get("image") else W - 2 * MX
        if s.get("image"):
            self.image(sl, s["image"], 7.55, 0, W - 7.55, H, "cover")
        y = 0.9
        fs = self.type["coverTitle"]
        while fs > self.type["title"] and self.lines(s["title"], fs, tw) > 3:
            fs -= 4
        th = self.block_h(s["title"], fs, tw, 1.12) + 0.1
        self.text(sl, s["title"], MX, y, tw, th, role="title", fs=fs, bold=True, color=t["white"], name="Baslik", spacing=1.12)
        y += th + 0.3
        if s.get("subtitle"):
            sh = self.block_h(s["subtitle"], self.type["coverSub"], tw) + 0.1
            self.text(sl, s["subtitle"], MX, y, tw, sh, fs=self.type["coverSub"], color="DCE3E8")
            y += sh + 0.3
        if s.get("meta"):
            self.text(sl, s["meta"], MX, max(y, 5.2), tw, self.block_h(s["meta"], self.type["body"], tw) + 0.1, color=t["white"], bold=True)
        self.frame(sl, {**s, "kicker": None}, dark=True, number=False, title=False)

    def agenda(self, s):
        sl = self.new(self.t["white"])
        items, tw = s.get("items", []), W - 2 * MX - 1.0
        hs = [max(0.72, self.block_h(i, self.type["body"], tw, 1.15)) for i in items]
        avail = G["bodyB"] - G["bodyY"]
        gap = max(0.1, min(0.3, (avail - sum(hs)) / max(len(items) - 1, 1)))
        y = G["bodyY"]
        for i, it in enumerate(items):
            self.chip(sl, i + 1, MX, y, 0.72, [self.t["accent"], self.t["accent2"], self.t["accent3"]][i % 3])
            self.text(sl, it, MX + 1.0, y, tw, hs[i], valign="middle")
            y += hs[i] + gap
        self.frame(sl, s)

    def section(self, s):
        t, sl = self.t, self.new(self.t["dark"])
        if s.get("num") is not None:
            self.text(sl, str(s["num"]), MX, 1.1, 3, 1.9, role="stat", fs=self.type["sectionNum"], bold=True, color=t["accent3"], name="Bolum no")
        self.text(sl, s["title"], MX, 3.15, W - 2 * MX, 1.6, role="title", fs=self.type["section"], bold=True, color=t["white"], name="Baslik")
        if s.get("subtitle"):
            self.text(sl, s["subtitle"], MX, 4.85, W - 2 * MX, 1.4, color="DCE3E8")
        self.frame(sl, s, dark=True, title=False)

    def bullets_layout(self, s):
        sl = self.new(self.t["white"])
        w = 6.6 if s.get("image") else W - 2 * MX
        self.bullets(sl, s.get("items", []), MX, G["bodyY"], w, G["bodyB"] - G["bodyY"])
        if s.get("image"):
            x = MX + w + 0.35
            self.image(sl, s["image"], x, G["bodyY"], W - MX - x, G["bodyB"] - G["bodyY"] - 0.35, s.get("imageMode", "contain"))
        self.frame(sl, s)

    def stat(self, s):
        t, sl = self.t, self.new(self.t["white"])
        st = s.get("stats", [])
        cw = (W - 2 * MX - 0.4 * (len(st) - 1)) / max(len(st), 1)
        for i, x in enumerate(st):
            cx = MX + i * (cw + 0.4)
            self.rect(sl, cx, G["bodyY"], cw, 3.55, t["light"], "Kart")
            self.text(sl, x["big"], cx + 0.3, G["bodyY"] + 0.25, cw - 0.6, 1.5, role="stat", fs=x.get("fs", self.type["stat"]), bold=True, color=[t["accent"], t["accent2"], t["accent3"]][i % 3], name="Buyuk sayi")
            self.text(sl, x["label"], cx + 0.3, G["bodyY"] + 1.8, cw - 0.6, 1.6, fs=self.type["statLabel"])
        if s.get("takeaway"):
            self.text(sl, s["takeaway"], MX, G["bodyY"] + 3.75, W - 2 * MX, 1.1, bold=True, color=t["dark"])
        self.frame(sl, s)

    def cards(self, s):
        t, sl = self.t, self.new(self.t["white"])
        cs = s.get("cards", [])
        gap, ch = 0.35, (3.75 if s.get("takeaway") else 4.8)
        cw = (W - 2 * MX - gap * (len(cs) - 1)) / max(len(cs), 1)
        for i, c in enumerate(cs):
            cx = MX + i * (cw + gap)
            col = [t["accent"], t["accent2"], t["accent3"], t["dark"]][i % 4]
            self.rect(sl, cx, G["bodyY"], cw, ch, t["light"], "Kart")
            off = 0.65 if c.get("image") else 0
            if c.get("image"):
                self.image(sl, c["image"], cx + 0.25, G["bodyY"] + 0.25, cw - 0.5, 1.5)
            else:
                self.chip(sl, c.get("mark", i + 1), cx + 0.3, G["bodyY"] + 0.3, 0.8, col)
            self.text(sl, c["head"], cx + 0.3, G["bodyY"] + 1.25 + off, cw - 0.6, 0.75, bold=True, color=col, name="Kart basligi")
            if c.get("body"):
                self.text(sl, c["body"], cx + 0.3, G["bodyY"] + 2.05 + off, cw - 0.6, ch - 2.2 - off)
        if s.get("takeaway"):
            self.text(sl, s["takeaway"], MX, G["bodyY"] + ch + 0.15, W - 2 * MX, 0.8, bold=True, color=t["dark"])
        self.frame(sl, s)

    def photoSteps(self, s):
        t, sl = self.t, self.new(self.t["white"])
        steps, lw = s.get("steps", []), s.get("textWidth", 7.4)
        tw, fs = lw - 0.85, self.type["body"]
        hs = [self.block_h(x, fs, tw, 1.15) for x in steps]
        avail = G["bodyB"] - G["bodyY"]
        gap = max(0.08, min(0.35, (avail - sum(hs)) / max(len(steps) - 1, 1)))
        if sum(hs) + gap * (len(steps) - 1) > avail + 0.05:
            self.warnings.append(f"slide {self.n}: steps too long — one line each or split")
        y = G["bodyY"]
        for i, st in enumerate(steps):
            self.chip(sl, i + 1, MX, y + 0.02, 0.62, t["accent"])
            self.text(sl, st, MX + 0.85, y, tw, hs[i] + 0.05, spacing=1.15)
            y += hs[i] + gap
        ix = MX + lw + 0.3
        iw = W - MX - ix
        ch = max(1.0, self.block_h(s["callout"], fs, iw - 0.4, 1.15) + 0.25) if s.get("callout") else 0
        self.image(sl, s["image"], ix, G["bodyY"], iw, G["bodyB"] - G["bodyY"] - ch - 0.1, s.get("imageMode", "cover"))
        if s.get("callout"):
            self.rect(sl, ix, G["bodyB"] - ch, iw, ch, t["light"], "Sonuc kutusu")
            self.text(sl, s["callout"], ix + 0.2, G["bodyB"] - ch + 0.08, iw - 0.4, ch - 0.16, align="center", valign="middle", bold=True, color=t["dark"], spacing=1.15)
        self.frame(sl, s)

    def compare(self, s):
        t, sl = self.t, self.new(self.t["white"])
        cw, ch = (W - 2 * MX - 0.5) / 2, (3.85 if s.get("takeaway") else 4.8)
        for i, (side, col) in enumerate([(s["left"], t["accent"]), (s["right"], t["accent2"])]):
            cx = MX + i * (cw + 0.5)
            self.rect(sl, cx, G["bodyY"], cw, ch, t["light"], "Kart")
            self.text(sl, side["head"], cx + 0.3, G["bodyY"] + 0.25, cw - 0.6, 0.75, bold=True, color=col, name="Kart basligi")
            off = 1.7 if side.get("image") else 0
            if side.get("image"):
                self.image(sl, side["image"], cx + 0.3, G["bodyY"] + 1.05, cw - 0.6, 1.6)
            self.bullets(sl, side.get("items", []), cx + 0.3, G["bodyY"] + 1.05 + off, cw - 0.6, ch - 1.2 - off)
        if s.get("takeaway"):
            self.text(sl, s["takeaway"], MX, G["bodyY"] + ch + 0.15, W - 2 * MX, 0.8, bold=True, color=t["dark"])
        self.frame(sl, s)

    def process(self, s):
        t, sl = self.t, self.new(self.t["white"])
        st = s.get("steps", [])
        gap = 0.45
        cw = (W - 2 * MX - gap * (len(st) - 1)) / max(len(st), 1)
        for i, x in enumerate(st):
            cx = MX + i * (cw + gap)
            self.chip(sl, i + 1, cx, G["bodyY"] + 0.1, 0.8, t["accent"])
            if i < len(st) - 1:
                ln = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(cx + 0.9), Inches(G["bodyY"] + 0.5), Inches(cx + cw + gap - 0.1), Inches(G["bodyY"] + 0.5))
                ln.line.color.rgb = rgb(t["muted"])
                ln.line.width = Pt(2.5)
                ln.name = "Ok"
                ln.line._get_or_add_ln().append(ln.line._get_or_add_ln().makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", {"type": "triangle"}))
            head = x.get("head", x) if isinstance(x, dict) else x
            self.text(sl, head, cx, G["bodyY"] + 1.15, cw, 0.75, bold=True, color=t["dark"])
            if isinstance(x, dict) and x.get("body"):
                self.text(sl, x["body"], cx, G["bodyY"] + 1.95, cw, 2.9)
        if s.get("takeaway"):
            self.text(sl, s["takeaway"], MX, 5.85, W - 2 * MX, 0.8, bold=True, color=t["dark"])
        self.frame(sl, s)

    def chart(self, s):
        t, sl = self.t, self.new(self.t["white"])
        cw = 7.6 if s.get("takeaway") else W - 2 * MX
        data = s.get("data", [])
        cd = CategoryChartData()
        cd.categories = data[0]["labels"] if data else []
        for ser in data:
            cd.add_series(ser.get("name", ""), ser["values"])
        kind = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT}.get(s.get("chartType", "bar"), XL_CHART_TYPE.COLUMN_CLUSTERED)
        if s.get("barDir") == "bar" and kind == XL_CHART_TYPE.COLUMN_CLUSTERED:
            kind = XL_CHART_TYPE.BAR_CLUSTERED
        gf = sl.shapes.add_chart(kind, Inches(MX), Inches(G["bodyY"]), Inches(cw), Inches(G["bodyB"] - G["bodyY"]), cd)
        gf.name = "Grafik"
        ch = gf.chart
        fs = Pt(s.get("labelFs", self.type["exceptions"].get("chart", self.type["minBody"])))
        ch.font.size = fs
        ch.font.name = self.font
        ch.has_title = False
        ch.has_legend = len(data) > 1
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
        plot = ch.plots[0]
        plot.has_data_labels = True
        plot.data_labels.font.size = fs
        cols = [t["accent"], t["accent2"], t["accent3"], t["dark"]]
        if kind in (XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT):
            for i, pt in enumerate(plot.series[0].points):
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = rgb(cols[i % 4])
        else:
            if len(data) == 1:
                plot.vary_by_categories = True
                for i, pt in enumerate(plot.series[0].points):
                    pt.format.fill.solid()
                    pt.format.fill.fore_color.rgb = rgb(cols[i % 4])
            else:
                for i, ser in enumerate(plot.series):
                    ser.format.fill.solid()
                    ser.format.fill.fore_color.rgb = rgb(cols[i % 4])
            ch.value_axis.has_major_gridlines = True
            ch.value_axis.major_gridlines.format.line.color.rgb = rgb("E3E7EB")
        if s.get("takeaway"):
            x = MX + cw + 0.35
            self.rect(sl, x, G["bodyY"], W - MX - x, G["bodyB"] - G["bodyY"], t["light"], "Yorum kutusu")
            self.text(sl, s["takeaway"], x + 0.25, G["bodyY"] + 0.3, W - MX - x - 0.5, G["bodyB"] - G["bodyY"] - 0.6, color=t["dark"])
        self.frame(sl, s)

    def table(self, s):
        t, sl = self.t, self.new(self.t["white"])
        rows = s.get("rows", [])
        fs = self.floor("body", s.get("fs"))
        if len(rows) > 5:
            self.warnings.append(f"slide {self.n}: table has {len(rows)} rows at {fs} pt — split or move to notes")
        if rows:
            gf = sl.shapes.add_table(len(rows), len(rows[0]), Inches(MX), Inches(G["bodyY"]), Inches(W - 2 * MX), Inches(min(4.8, 0.8 * len(rows))))
            gf.name = "Tablo"
            tb = gf.table
            for ri, r in enumerate(rows):
                for ci, c in enumerate(r):
                    cell = tb.cell(ri, ci)
                    cell.text = ""
                    self._runs(cell.text_frame.paragraphs[0], str(c), t["white"] if ri == 0 else t["ink"], ri == 0, fs)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = rgb(t["dark"] if ri == 0 else (t["light"] if ri % 2 else t["white"]))
        self.frame(sl, s)

    def quote(self, s):
        dark = s.get("dark", True) is not False
        t, sl = self.t, self.new(self.t["dark"] if dark else self.t["light"])
        self.text(sl, s["text"], 1.1, 1.6, W - 2.2, 3.4, fs=s.get("fs", 44), bold=True, color=t["white"] if dark else t["dark"], valign="middle")
        if s.get("by"):
            self.text(sl, s["by"], 1.1, 5.15, W - 2.2, 0.9, color=t["accent3"] if dark else t["accent"])
        self.frame(sl, s, dark=dark)

    def references(self, s):
        t, sl = self.t, self.new(self.t["white"])
        fs = self.type["exceptions"]["references"]
        items = s.get("items", [])
        tb = sl.shapes.add_textbox(Inches(MX), Inches(G["bodyY"] - 0.2), Inches(W - 2 * MX), Inches(G["bodyB"] - G["bodyY"] + 0.6))
        tb.name = "Kaynakca"
        tf = tb.text_frame
        tf.word_wrap = True
        for i, r in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(4)
            self._runs(p, r, t["ink"], False, fs)
        if self.block_h("\n".join(items), fs, W - 2 * MX, 1.25) > 5.2:
            self.warnings.append(f"slide {self.n}: references overflow — split into another references slide")
        self.frame(sl, {**s, "source": None})

    def closing(self, s):
        t, sl = self.t, self.new(self.t["dark"])
        self.text(sl, s.get("title", "Özet"), MX, 0.6, W - 2 * MX, 1.0, role="title", fs=self.type["title"], bold=True, color=t["white"], name="Baslik")
        self.bullets(sl, s.get("items", []), MX, 1.9, W - 2 * MX, 4.6, color=t["white"])
        self.frame(sl, s, dark=True, title=False)

    LAYOUTS = ["cover", "agenda", "section", "bullets", "stat", "cards", "photoSteps", "compare", "process", "chart", "table", "quote", "references", "closing"]

    def add(self, s):
        self.n += 1
        lay = s["layout"]
        if lay not in self.LAYOUTS:
            raise SystemExit(f"slide {self.n}: unknown layout {lay}; use one of {self.LAYOUTS}")
        self.log.append(lay)
        if len(self.log) >= 3 and len(set(self.log[-3:])) == 1 and lay != "references":
            self.warnings.append(f"slide {self.n}: layout {lay} 3 times in a row — vary the layout")
        getattr(self, "bullets_layout" if lay == "bullets" else lay)(s)


# ---------------------------------------------------------------- notes (same rules as build.js)
def build_notes(deck):
    meta = deck.get("meta", {})
    wpm, pause = meta.get("wordsPerMinute", 112), meta.get("pausePerSlide", 7)
    mode = meta.get("abbreviationMode", "every")
    abbr, names = deck.get("abbreviations", {}), deck.get("names", {})
    soft, hard = meta.get("noteWordsSoft", 120), meta.get("noteWordsHard", 160)
    seen: set = set()
    warns: list[str] = []

    def expand(t):
        if not t:
            return t
        for ab, (say, full) in abbr.items():
            pat = re.compile(r"(?<![\w\-])" + re.escape(ab) + r"((?:['’]\w+)?)(?![\w])(?! \()")
            k = [0]

            def rep(m):
                k[0] += 1
                return f"{ab}{m.group(1)} ({say}; {full})" if (mode == "every" or k[0] == 1) else f"{ab}{m.group(1)} ({say})"
            t = pat.sub(rep, t)
        for nm, say in names.items():
            if nm in seen:
                continue
            pat = re.compile(r"(?<!\w)" + re.escape(nm) + r"((?:['’]\w+)?)(?!\w)")
            if pat.search(t):
                t = pat.sub(lambda m: f"{nm}{m.group(1)} ({say})", t, count=1)
                seen.add(nm)
        return t

    def words(t):
        return len(re.sub(r"\([^)]*\)", " ", t or "").split())

    def r5(x):
        return max(5, round(x / 5) * 5)

    bodies = []
    for i, s in enumerate(deck["slides"], 1):
        n = s.get("notes")
        if s["layout"] == "references":
            bodies.append(("", 0, 5))
            continue
        if not n:
            bodies.append(("", 0, 10 if s["layout"] == "section" else pause))
            warns.append(f"slide {i}: no speaker notes")
            continue
        if isinstance(n, str):
            w = words(n)
            bodies.append((expand(n), w, r5(w / wpm * 60 + pause)))
            continue
        parts, spoken, half = [], words(n.get("say")) + words(n.get("summary")) + words(n.get("transition")), 0
        if n.get("say"):
            parts.append(expand(n["say"]))
        if n.get("summary"):
            parts.append("ÖZETLE: " + expand(n["summary"]))
        rs = n.get("research") or []
        rs = rs if isinstance(rs, list) else [rs]
        for r in rs:
            if isinstance(r, str):
                parts.append("ARAŞTIRMA: " + expand(r))
                spoken += words(r)
                continue
            lines = [("ARAŞTIRMA — " + r.get("cite", "")).strip()]
            if r.get("why"):
                lines.append("• Neden burada: " + expand(r["why"]))
                half += words(r["why"])
            if r.get("what"):
                lines.append("• Ne yapılmış, ne bulunmuş: " + expand(r["what"]))
                spoken += words(r["what"])
            if r.get("interpret"):
                lines.append("• Nasıl yorumlanır: " + expand(r["interpret"]))
                spoken += words(r["interpret"])
            elif r.get("what") and re.search(r"\d", r["what"]):
                warns.append(f"slide {i}: research numbers without interpret — add threshold + meaning")
            parts.append("\n".join(lines))
        if n.get("terms"):
            tl = []
            for t in n["terms"]:
                if isinstance(t, str):
                    tl.append("• " + t)
                    continue
                ab = abbr.get(t["term"])
                say = t.get("say") or (ab[0] if ab else None)
                full = t.get("full") or (ab[1] if ab else "")
                head = f"{t['term']} ({say})" if say else t["term"]
                bits = [full, f"({t['english']})" if t.get("english") else "", ("— " if full else "") + expand(t["def"]) if t.get("def") else ""]
                tl.append(f"• {head}: " + " ".join(b for b in bits if b))
            parts.append("TERİMLER:\n" + "\n".join(tl))
        if n.get("question"):
            parts.append("OLASI SORU: " + expand(n["question"]))
        if n.get("transition"):
            parts.append("→ Geçiş: " + expand(n["transition"]))
        sw = spoken + 0.5 * half
        if sw > hard:
            warns.append(f"slide {i}: notes {round(sw)} spoken words > {hard}")
        elif sw > soft:
            warns.append(f"slide {i}: notes {round(sw)} spoken words > {soft} (soft)")
        bodies.append(("\n\n".join(parts), round(sw), r5(sw / wpm * 60 + pause)))
    total = sum(b[2] for b in bodies)
    fmt = lambda x: f"{int(x // 60)}:{int(round(x % 60)):02d}"  # noqa: E731
    acc, out = 0, []
    N = len(deck["slides"])
    for i, (body, w, sec) in enumerate(bodies, 1):
        acc += sec
        head = "" if meta.get("noteHeader") is False else f"[{i}/{N}]  Hedef ≈ {fmt(sec)} · Bitişte ≈ {fmt(acc)} / {fmt(total)}"
        out.append(("\n\n".join(x for x in (head, body) if x), w, sec))
    return out, total, warns, fmt


def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:  # noqa: BLE001
        pass
    src = sys.argv[1] if len(sys.argv) > 1 else "deck.json"
    deck = json.load(open(src, encoding="utf-8"))
    meta = deck.get("meta", {})
    out = sys.argv[2] if len(sys.argv) > 2 else meta.get("output", "sunum.pptx")
    notes, total, warns, fmt = build_notes(deck)
    d = Deck(meta, os.path.dirname(os.path.abspath(src)))
    forbid = [f.lower() for f in meta.get("forbiddenOnSlides", ["bu akşam", "bugün", "iyi akşamlar", "geçen hafta", "haftaya", "tartışalım", "öğretim örneği", "sorular?"])]
    for i, s in enumerate(deck["slides"], 1):
        content = s["layout"] not in ("cover", "section", "references", "closing", "agenda")
        if content and not s.get("source") and not s.get("noSource") and meta.get("sourceEverySlide", True):
            warns.append(f"slide {i}: content slide without source line")
        on = json.dumps({k: v for k, v in s.items() if k != "notes"}, ensure_ascii=False).lower()
        warns += [f'slide {i}: forbidden on-slide phrase "{f}"' for f in forbid if f in on]
        d.add({**s, "notes": notes[i - 1][0]})
    d.prs.save(out)
    print(f"Saved {out} — {d.n} slides, estimated talk {fmt(total)}")
    tgt = meta.get("targetMinutes")
    if tgt and not tgt[0] <= total / 60 <= tgt[1]:
        warns.append(f"timing {fmt(total)} outside target {tgt[0]}–{tgt[1]} min — adjust notes/slide count, never shrink type")
    for i, (s, (_, w, sec)) in enumerate(zip(deck["slides"], notes), 1):
        print(f"{i:2}  {s['layout']:<12} {str(s.get('title', ''))[:46]:<46} {fmt(sec):>5} {w:>5}")
    allw = d.warnings + warns
    if allw:
        print(f"\n{len(allw)} warning(s):")
        for w in allw:
            print(" -", w)
        return 2
    print("\nNo build warnings. Next: render + visual QA + pptx_audit.py")
    return 0


if __name__ == "__main__":
    sys.exit(main())
